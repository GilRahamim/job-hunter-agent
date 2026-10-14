"""CV file -> CandidateProfile (LLM extraction with a heuristic fallback, cached by file hash)."""
from __future__ import annotations

import hashlib
import logging
import re
from pathlib import Path
from typing import Any

from jobhunter import db
from jobhunter.config import ROOT, load_config
from jobhunter.models import CandidateProfile
from jobhunter.prompts import PROFILE_SYSTEM, PROFILE_USER
from jobhunter.skills import extract_skills

log = logging.getLogger(__name__)

CV_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}
MIN_PDF_CHARS = 200
HEURISTIC_SUFFIX = ":heuristic"
NOT_A_CV = {"cv_master", "template"}      # files the CV Tailoring Agent keeps in cv/


class CVError(RuntimeError):
    """The CV is missing or unreadable; the message tells the owner what to do."""


# ------------------------------------------------------------ reading

def find_cv(cv_dir: str | Path) -> Path:
    folder = Path(cv_dir)
    if not folder.is_absolute():
        folder = ROOT / folder
    files = [p for p in folder.glob("*")
             if p.suffix.lower() in CV_EXTENSIONS and p.name.lower() != "readme.md"
             and not p.name.startswith("~$") and p.stem.lower() not in NOT_A_CV]
    if not files:
        raise CVError(f"No CV found in {folder}. Put a PDF/DOCX/TXT file there.")
    return max(files, key=lambda p: p.stat().st_mtime)      # newest wins


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    text = "\n".join((page.extract_text() or "") for page in PdfReader(path).pages)
    if len(text.strip()) < MIN_PDF_CHARS:
        raise CVError(f"{path.name} has almost no extractable text (scanned PDF?). "
                      "Export the CV as a text PDF or DOCX.")
    return text


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:                  # many CV templates keep skills in tables
        for row in table.rows:
            seen: set[int] = set()
            for cell in row.cells:
                if id(cell._tc) not in seen:  # merged cells repeat
                    seen.add(id(cell._tc))
                    parts.append(cell.text)
    return "\n".join(parts)


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        raw = _read_pdf(path)
    elif suffix == ".docx":
        raw = _read_docx(path)
    else:
        raw = path.read_text(encoding="utf-8", errors="replace")
    lines = (re.sub(r"[ \t\xa0]+", " ", line).strip() for line in raw.splitlines())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def file_hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


# ------------------------------------------------------------ extraction

def extract_with_llm(llm: Any, cv_text: str) -> CandidateProfile:
    structured = llm.with_structured_output(CandidateProfile)
    return structured.invoke([("system", PROFILE_SYSTEM),
                              ("user", PROFILE_USER.format(cv_text=cv_text))])


def extract_heuristic(cv_text: str, cfg: dict[str, Any]) -> CandidateProfile:
    lines = [ln for ln in cv_text.splitlines() if ln.strip()]
    return CandidateProfile(
        name=lines[0] if lines else "",
        headline=" ".join(lines[:2]),
        summary=" ".join(lines[:2]),
        years_experience=0,
        seniority="junior",
        skills=extract_skills(cv_text),
        target_roles=list(cfg.get("search", {}).get("queries", [])),
    )


def apply_overrides(profile: CandidateProfile, cfg: dict[str, Any]) -> CandidateProfile:
    over = (cfg.get("profile") or {}).get("overrides") or {}
    skills = list(profile.skills)
    for skill in over.get("add_skills") or []:
        if skill.lower() not in {s.lower() for s in skills}:
            skills.append(skill)
    remove = {s.lower() for s in over.get("remove_skills") or []}
    skills = [s for s in skills if s.lower() not in remove]
    data = profile.model_copy(update={"skills": skills})
    if over.get("seniority"):
        data = data.model_copy(update={"seniority": over["seniority"]})
    return data


def load_profile(cfg: dict[str, Any], llm: Any | None, db_path: str | Path | None = None,
                 use_cache: bool = True) -> tuple[CandidateProfile, str, bool]:
    """Returns (profile, cv_hash, cache_hit). `llm=None` -> heuristic profile."""
    path = find_cv(cfg["profile"]["cv_dir"])
    cv_hash = file_hash(path)

    if use_cache:
        keys = [cv_hash] if llm else [cv_hash, cv_hash + HEURISTIC_SUFFIX]
        for key in keys:
            cached = db.load_profile(key, db_path)
            if cached:
                log.info("profile cache hit (%s)", key)
                return apply_overrides(CandidateProfile(**cached), cfg), cv_hash, True

    text = read_text(path)
    profile: CandidateProfile | None = None
    key = cv_hash
    if llm:
        try:
            profile = extract_with_llm(llm, text)
        except Exception as e:  # noqa: BLE001 — a failed call must not block the run
            log.warning("LLM profile extraction failed (%s); using heuristic profile", e)
    if profile is None:
        profile, key = extract_heuristic(text, cfg), cv_hash + HEURISTIC_SUFFIX
    if use_cache:
        db.save_profile(key, path.name, profile.model_dump(), db_path)
    return apply_overrides(profile, cfg), cv_hash, False


def main() -> None:
    import argparse
    import sys

    from jobhunter.llm import get_llm

    parser = argparse.ArgumentParser(description="Print the extracted candidate profile")
    parser.add_argument("--no-llm", action="store_true", help="heuristic profile only")
    args = parser.parse_args()
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")

    cfg = load_config()
    llm = None if args.no_llm else get_llm(cfg)
    try:
        profile, cv_hash, hit = load_profile(cfg, llm, use_cache=not args.no_llm)
    except CVError as e:
        sys.exit(str(e))
    print(profile.model_dump_json(indent=2))
    print(f"cv_hash: {cv_hash}\ncache: {'hit' if hit else 'miss'}\n"
          f"extractor: {'llm' if llm and db.load_profile(cv_hash) else 'heuristic'}")


if __name__ == "__main__":
    main()
