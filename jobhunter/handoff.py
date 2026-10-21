"""Hand one job to Claude on the person's subscription (Claude Code or a claude.ai chat), not the API.

The search finds and scores jobs; for the one you want to apply to, Claude does the writing:
a tailored CV, a cover letter, interview prep and a list of what changed. This module writes
<data>/tailored/<job>/brief.md with the posting, the search agent's analysis, the CV and the rules
(prompts.HANDOFF_BRIEF), and can open Claude Code on it. Claude saves its files next to the brief.

  python -m jobhunter.handoff --list                    # best jobs and their ids
  python -m jobhunter.handoff <job_id>                  # write the brief, print how to start Claude
  python -m jobhunter.handoff <job_id> --launch         # ... and open Claude Code on it
  python -m jobhunter.handoff <job_id> --headless       # ... and let Claude write the kit, no window
  python -m jobhunter.handoff --docx <folder>/cv_tailored.md   # Markdown -> Word
  python -m jobhunter.handoff --apply-edits <folder>/cv_edits.json   # owner: edit a copy of the original .docx
  (add --user <name> for a user added in the dashboard)
"""
from __future__ import annotations

import argparse
import json
import os
import re
import shlex
import shutil
import subprocess
import sys
import time
from copy import deepcopy
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from jobhunter import db
from jobhunter.config import ROOT, active_user, load_config
from jobhunter.cv import CVError, find_cv, read_text
from jobhunter.matcher import explanation_language
from jobhunter.prompts import HANDOFF_BRIEF, HANDOFF_DOCX_NOTE, HANDOFF_FILES_NOTE, HANDOFF_OWNER_BLOCK

BRIEF = "brief.md"
EDITS = "cv_edits.json"
HEADLESS_LOG = "auto_tailor.json"
OWNER_RULES = ROOT / ".claude" / "skills" / "cv-tailoring" / "references" / "owner-cv-rules.md"
MAX_DESCRIPTION_CHARS = 15000
# A headless run that hits the plan's usage window says so in its result text.
_LIMIT = re.compile(r"usage limit|session limit|rate limit|limit reached|\b429\b|too many requests", re.I)
# Claude Code ranks these above the subscription login (an approved ANTHROPIC_API_KEY bills the API),
# and the dashboard has .env loaded: drop them so the launched session runs on the Claude subscription.
API_KEY_VARS = ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN")
_SLUG = re.compile(r"[^a-z0-9]+")
_HEBREW = re.compile(r"[֐-׿]")


class HandoffError(RuntimeError):
    """Something the person has to fix (unknown job, Claude Code not installed)."""


# ------------------------------------------------------------ paths

def tailored_dir(db_file: str | Path) -> Path:
    """Next to the user's jobs.db: data/tailored for the owner, users/<id>/data/tailored for others."""
    return Path(db_file).parent / "tailored"


def _slug(text: str, n: int) -> str:
    return _SLUG.sub("-", (text or "").lower()).strip("-")[:n].strip("-")


def job_folder(job: dict[str, Any], db_file: str | Path) -> Path:
    """ASCII folder name ("acme-junior-python-developer-3f9a1c"), safe on Windows command lines."""
    name = "-".join(p for p in (_slug(job["company"], 20), _slug(job["title"], 40), job["id"][:6]) if p)
    return tailored_dir(db_file) / name


def kit_files(folder: Path) -> list[Path]:
    """What Claude wrote for this job (everything but the brief), newest first."""
    if not folder.is_dir():
        return []
    files = [p for p in folder.iterdir() if p.is_file() and p.name not in (BRIEF, HEADLESS_LOG)]
    return sorted(files, key=lambda p: p.stat().st_mtime, reverse=True)


# ------------------------------------------------------------ the brief

def _language_of(text: str) -> str:
    letters = re.findall(r"[^\W\d_]", text)
    hebrew = sum(1 for ch in letters if _HEBREW.match(ch))
    return "Hebrew" if letters and hebrew / len(letters) > 0.3 else "English"


def _analysis(job: dict[str, Any]) -> str:
    score = f"Match score: {job['score'] or 0:.0f}/100"
    if not job.get("verdict"):
        return f"{score} (keyword pre-score only; no LLM review of this job yet)."
    lines = [f"{score}, verdict: {job['verdict']} ({job.get('scored_by') or ''})"]
    if job.get("reasons"):
        lines.append(f"Why: {job['reasons']}")
    if job.get("seniority_fit"):
        lines.append(f"Seniority fit: {job['seniority_fit']}")
    if job.get("matching_skills"):
        lines.append(f"Matching skills: {', '.join(job['matching_skills'])}")
    if job.get("missing_skills"):
        lines.append(f"Missing skills: {', '.join(job['missing_skills'])}")
    if job.get("cv_tip"):
        lines.append(f"CV tip: {job['cv_tip']}")
    return "\n".join(lines)


def _cv(cfg: dict[str, Any]) -> tuple[str, str, str]:
    """(CV file path, CV text, master CV block) for the brief."""
    try:
        path = find_cv(cfg["profile"]["cv_dir"])
        cv_file, cv_text = str(path), read_text(path)
    except CVError as e:
        return "(none)", f"(the CV could not be read: {e})", ""
    master = path.parent / "cv_master.yaml"
    block = ""
    if master.exists():
        block = ("\n## Master CV (the richer source; the same rule applies)\n\n```yaml\n"
                 f"{master.read_text(encoding='utf-8').strip()}\n```\n")
    return cv_file, cv_text, block


def _command(*args: str) -> str:
    """A handoff command with this Python (it has python-docx), so it works without an activated venv."""
    parts = [sys.executable, "-m", "jobhunter.handoff", *args]
    return subprocess.list2cmdline(parts) if os.name == "nt" else shlex.join(parts)


def docx_command(md_file: Path) -> str:
    return _command("--docx", str(md_file))


# ------------------------------------------------------------ the owner's CV (structure lock)

def original_docx(cv_dir: str | Path) -> Path | None:
    """The newest Word CV in cv/: the base every tailored CV of the owner is edited from."""
    folder = Path(cv_dir) if Path(cv_dir).is_absolute() else ROOT / cv_dir
    files = [p for p in folder.glob("*.docx") if not p.name.startswith("~$") and p.stem.lower() != "template"]
    return max(files, key=lambda p: p.stat().st_mtime) if files else None


def _locked(index: int, paragraph: Any) -> bool:
    """Paragraphs the structure lock forbids touching: name, contact, headings, entry headers, blanks."""
    # TODO(human)
    raise NotImplementedError


def paragraph_listing(docx_file: Path) -> str:
    from docx import Document

    lines = []
    for i, p in enumerate(Document(docx_file).paragraphs):
        mark = "[locked] " if _locked(i, p) else ""
        lines.append(f"{i:>3} {mark}{p.text.replace(chr(9), ' <tab> ')}")
    return "\n".join(lines)


def _note(path: Path) -> str:
    return path.read_text(encoding="utf-8").strip() if path.exists() else "(not available)"


def owner_block(cfg: dict[str, Any], folder: Path) -> str:
    """The owner's rules, notes and paragraph map; empty when there is no original .docx to edit."""
    original = original_docx(cfg["profile"]["cv_dir"])
    if not original:
        return ""
    notes = original.parent / "notes"
    return HANDOFF_OWNER_BLOCK.format(
        rules_file=OWNER_RULES.relative_to(ROOT).as_posix(), original=original,
        paragraphs=paragraph_listing(original), command=_command("--apply-edits", str(folder / EDITS)),
        master_cv=_note(notes / "master_cv.md"), experience_bank=_note(notes / "experience_bank.md"))


def _set_text(paragraph: Any, text: str) -> None:
    """Replace a paragraph's text, keeping its run formatting; **x** marks bold."""
    runs = paragraph.runs
    bold_src = next((r for r in runs if r.bold), None)
    plain_src = next((r for r in runs if not r.bold), None)
    for r in runs:
        r._r.getparent().remove(r._r)
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        bold = bool(i % 2)
        src = (bold_src if bold else plain_src) or plain_src or bold_src
        run = paragraph.add_run(part)
        if src is not None and src._r.rPr is not None:
            run._r.insert(0, deepcopy(src._r.rPr))
        run.bold = bold


def apply_docx_edits(edits_file: str | Path, base: str | Path | None = None) -> tuple[Path, list[str]]:
    """Copy the original .docx and replace the text of the paragraphs named in cv_edits.json.

    Returns the new file and warnings: locked or unknown paragraphs are skipped, and a paragraph
    much longer than before may wrap and push the CV to a second page."""
    from docx import Document

    edits_file = Path(edits_file)
    spec = json.loads(edits_file.read_text(encoding="utf-8"))
    base = Path(base) if base else original_docx(load_config("")["profile"]["cv_dir"])
    if not base or not base.exists():
        raise HandoffError(f"no original .docx to edit ({base})")
    doc = Document(base)
    paragraphs = doc.paragraphs
    warnings = []
    for key, text in spec.get("edits", {}).items():
        i = int(key)
        if not 0 <= i < len(paragraphs) or _locked(i, paragraphs[i]):
            warnings.append(f"paragraph {i}: locked or unknown, skipped")
            continue
        old = len(paragraphs[i].text)
        _set_text(paragraphs[i], text)
        new = len(paragraphs[i].text)
        if new > old + max(8, old * 0.03):
            warnings.append(f"paragraph {i}: {old} -> {new} characters, may wrap; shorten it")
    out = edits_file.parent / Path(spec.get("out") or "cv_tailored.docx").name
    doc.save(out)
    return out, warnings


def build_brief(job: dict[str, Any], cfg: dict[str, Any], profile: dict[str, Any] | None,
                folder: Path, owner: bool = False) -> str:
    cv_file, cv_text, master_block = _cv(cfg)
    name = (profile or {}).get("name") or cfg.get("display_name") or "the candidate"
    description = (job.get("description") or "(no description was collected; open the link)").strip()
    return HANDOFF_BRIEF.format(
        name=name, title=job["title"], company=job["company"],
        location=job.get("location") or job.get("remote_region") or ("Remote" if job.get("is_remote") else "-"),
        url=job.get("url") or "-", description=description[:MAX_DESCRIPTION_CHARS],
        posting_language=_language_of(f"{job['title']} {description}"),
        language=explanation_language(cfg),
        output_note=HANDOFF_FILES_NOTE.format(folder=folder),
        docx_note=HANDOFF_DOCX_NOTE.format(command=docx_command(folder / "cv_tailored.md")),
        analysis=_analysis(job), cv_file=cv_file, cv_text=cv_text,
        master_block=master_block + (owner_block(cfg, folder) if owner else ""))


def write_brief(job_id: str, user: str | None = None, db_file: str | Path | None = None) -> Path:
    """Write (or refresh) the brief for one stored job. Returns its path."""
    user = active_user() if user is None else user
    db_file = db.db_path(db_file, user)
    job = db.job_detail(job_id, db_file)
    if not job:
        raise HandoffError(f"no job with id {job_id!r} (python -m jobhunter.handoff --list shows the ids)")
    cfg = load_config(user)
    folder = job_folder(job, db_file)
    folder.mkdir(parents=True, exist_ok=True)
    latest = db.latest_profile(db_file)
    brief = folder / BRIEF
    brief.write_text(build_brief(job, cfg, latest["data"] if latest else None, folder, owner=not user),
                     encoding="utf-8")
    return brief


# ------------------------------------------------------------ starting Claude

def claude_prompt(brief: Path) -> str:
    """Short on purpose (Windows command lines): everything else is in the brief."""
    try:
        where = brief.resolve().relative_to(ROOT).as_posix()
    except ValueError:
        where = brief.resolve().as_posix()
    return f"Read {where} and follow its instructions."


def shell_command(brief: Path) -> str:
    """What to type in a terminal at the project folder."""
    return f'claude "{claude_prompt(brief)}"'


def subscription_env(env: dict[str, str] | None = None) -> dict[str, str]:
    env = dict(os.environ if env is None else env)
    for var in API_KEY_VARS:
        env.pop(var, None)
    return env


def claude_path() -> str | None:
    return shutil.which("claude")


def _system() -> str:
    return "windows" if os.name == "nt" else sys.platform


def can_launch(system: str | None = None) -> bool:
    """A new terminal window can be opened here (Windows or macOS) and Claude Code is installed."""
    return bool(claude_path()) and (system or _system()) in ("windows", "darwin")


def launch_claude(brief: Path, system: str | None = None) -> None:
    """Open Claude Code in a new terminal window on the brief, signed in with the subscription."""
    exe = claude_path()
    if not exe:
        raise HandoffError("Claude Code is not installed (see https://code.claude.com/docs); "
                           "or attach brief.md to a chat on claude.ai")
    prompt = claude_prompt(brief)
    system = system or _system()
    if system == "windows":
        subprocess.Popen([exe, prompt], cwd=ROOT, env=subscription_env(),
                         creationflags=getattr(subprocess, "CREATE_NEW_CONSOLE", 0))
    elif system == "darwin":
        line = (f"cd {shlex.quote(str(ROOT))} && unset {' '.join(API_KEY_VARS)} && "
                f"claude {shlex.quote(prompt)}")
        line = line.replace("\\", "\\\\").replace('"', '\\"')
        subprocess.Popen(["osascript", "-e", f'tell application "Terminal" to do script "{line}"',
                          "-e", 'tell application "Terminal" to activate'])
    else:
        raise HandoffError(f"open a terminal in {ROOT} and run: {shell_command(brief)}")


# ------------------------------------------------------------ headless (no window, subscription)

@dataclass
class HeadlessResult:
    job_id: str
    ok: bool
    result: str                          # Claude's final message, or why nothing ran
    folder: Path | None = None
    exit_code: int | None = None
    limit_hit: bool = False              # the plan's usage window is used up: stop the queue
    session_id: str | None = None
    cost_usd: float | None = None        # client-side estimate; the subscription bills nothing per call
    duration_s: float = 0.0


def allowed_tools() -> list[str]:
    """The only shell command /apply needs, in the spellings Claude may type it."""
    prefixes = [_command(), "python -m jobhunter.handoff",
                ".venv/Scripts/python.exe -m jobhunter.handoff", r".venv\Scripts\python.exe -m jobhunter.handoff"]
    return [f"Bash({p} *)" for p in dict.fromkeys(prefixes)]


def headless_command(exe: str, job_id: str, user: str = "") -> list[str]:
    prompt = f"/apply {job_id}" + (f" --user {user}" if user else "") + " --unattended"
    return [exe, "-p", prompt,                          # never --bare: it skips the login and the skills
            "--permission-mode", "acceptEdits", "--permission-prompts", "none",
            "--allowedTools", *allowed_tools(), "--output-format", "json"]


def _finish_kit(folder: Path, base: Path | None) -> None:
    """The Word copy, in Python, when Claude's shell command was not allowed to run."""
    edits = folder / EDITS
    if edits.exists() and base:
        out = json.loads(edits.read_text(encoding="utf-8")).get("out") or "cv_tailored.docx"
        if not (folder / Path(out).name).exists():
            apply_docx_edits(edits, base)
    elif (folder / "cv_tailored.md").exists() and not (folder / "cv_tailored.docx").exists():
        markdown_to_docx(folder / "cv_tailored.md")


def _parse(stdout: str) -> dict[str, Any]:
    try:
        data = json.loads(stdout)
        return data if isinstance(data, dict) else {}
    except (json.JSONDecodeError, TypeError):
        return {}


def _log(res: HeadlessResult) -> None:
    if res.folder and res.folder.is_dir():
        (res.folder / HEADLESS_LOG).write_text(json.dumps({
            "exit_code": res.exit_code, "ok": res.ok, "limit_hit": res.limit_hit,
            "session_id": res.session_id, "result": res.result[:500], "total_cost_usd": res.cost_usd,
            "duration_s": round(res.duration_s, 1),
            "at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        }, ensure_ascii=False, indent=2), encoding="utf-8")


def run_headless(job_id: str, user: str | None = None, db_file: str | Path | None = None,
                 timeout: int = 900) -> HeadlessResult:
    """Let Claude Code write one job's kit with no window, on the subscription. Never raises."""
    user = active_user() if user is None else user
    res = HeadlessResult(job_id, ok=False, result="")
    start = time.monotonic()
    try:
        exe = claude_path()
        if not exe:
            res.result = "Claude Code is not installed"
            return res
        res.folder = write_brief(job_id, user, db_file).parent
        proc = subprocess.run(headless_command(exe, job_id, user), cwd=ROOT, env=subscription_env(),
                              capture_output=True, text=True, encoding="utf-8", errors="replace",
                              timeout=timeout)
        data = _parse(proc.stdout)
        res.exit_code = proc.returncode
        res.result = str(data.get("result") or proc.stdout or proc.stderr or "").strip()
        res.session_id = data.get("session_id")
        res.cost_usd = data.get("total_cost_usd")
        res.ok = proc.returncode == 0 and not data.get("is_error")
        res.limit_hit = not res.ok and bool(_LIMIT.search(f"{res.result} {proc.stderr or ''}"))
        if res.ok:
            base = None if user else original_docx(load_config(user)["profile"]["cv_dir"])
            _finish_kit(res.folder, base)
    except subprocess.TimeoutExpired:
        res.result = f"timed out after {timeout}s"
    except Exception as e:                                   # one job must never stop a run
        res.result = f"{type(e).__name__}: {e}"
    finally:
        res.duration_s = time.monotonic() - start
        try:
            _log(res)
        except OSError:
            pass
    return res


# ------------------------------------------------------------ Markdown -> Word

def _add_runs(paragraph: Any, text: str) -> None:
    """**bold** inside a line; everything else plain."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            paragraph.add_run(part).bold = bool(i % 2)


def _rtl(paragraph: Any) -> None:
    """Right-to-left paragraph for Hebrew lines."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement

    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:bidi"))


def markdown_to_docx(md_file: str | Path, docx_file: str | Path | None = None) -> Path:
    """A plain Word copy of a Markdown CV or letter: headings, bullets, **bold**, paragraphs."""
    from docx import Document
    from docx.shared import Pt

    md_file = Path(md_file)
    out = Path(docx_file) if docx_file else md_file.with_suffix(".docx")
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    for raw in md_file.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or re.fullmatch(r"[-*_]{3,}", line):
            continue
        heading = re.match(r"(#{1,3})\s+(.*)", line)
        bullet = re.match(r"[-*•]\s+(.*)", line)
        if heading:
            p = doc.add_heading(heading.group(2).replace("**", ""), level=len(heading.group(1)))
        elif bullet:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, bullet.group(1))
        else:
            p = doc.add_paragraph()
            _add_runs(p, line)
        if _HEBREW.search(line):
            _rtl(p)
    doc.save(out)
    return out


# ------------------------------------------------------------ CLI

def _list(user: str, n: int = 20) -> None:
    rows = [r for r in db.jobs_table(db.db_path(user=user)) if r["status"] not in ("rejected", "not_relevant")]
    rows.sort(key=lambda r: r["score"] or 0, reverse=True)
    for r in rows[:n]:
        print(f"{r['score'] or 0:5.1f}  {r['id']}  {r['title']} | {r['company']} | {r['status']}")
    if not rows:
        print("no jobs yet: run python run_agent.py first")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Prepare one job for Claude (subscription, not API)")
    parser.add_argument("job_id", nargs="?", help="id of a stored job (see --list)")
    parser.add_argument("--list", action="store_true", help="show the best jobs and their ids")
    parser.add_argument("--launch", action="store_true", help="open Claude Code on the brief (Windows/macOS)")
    parser.add_argument("--headless", action="store_true",
                        help="let Claude Code write the kit with no window (subscription, never the API)")
    parser.add_argument("--timeout", type=int, default=900, help="seconds for --headless (default 900)")
    parser.add_argument("--docx", metavar="MD_FILE", help="convert a Markdown file to Word and exit")
    parser.add_argument("--apply-edits", metavar="EDITS_JSON",
                        help="owner: copy the original .docx in cv/ and apply cv_edits.json to it")
    parser.add_argument("--user", default=None, help="a user added in the dashboard (default: the owner)")
    args = parser.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    user = args.user if args.user is not None else active_user()

    if args.docx:
        print(markdown_to_docx(args.docx))
        return 0
    if args.apply_edits:
        try:
            out, warnings = apply_docx_edits(args.apply_edits)
        except (HandoffError, OSError, ValueError, KeyError) as e:
            print(f"error: {e}")
            return 1
        print("\n".join([str(out), *(f"warning: {w}" for w in warnings)]))
        return 0
    if args.list or not args.job_id:
        _list(user)
        return 0
    if args.headless:
        res = run_headless(args.job_id, user, timeout=args.timeout)
        print(f"{'ok' if res.ok else 'failed'} in {res.duration_s:.0f}s  folder: {res.folder}")
        if res.folder:
            print("kit: " + ", ".join(p.name for p in kit_files(res.folder)))
        print(res.result)
        return 0 if res.ok else 1
    try:
        brief = write_brief(args.job_id, user)
        print(f"brief: {brief}\nstart Claude in {ROOT}:\n  {shell_command(brief)}\n"
              "or attach the brief to a new chat on claude.ai")
        if args.launch:
            launch_claude(brief)
            print("Claude Code opened in a new window.")
    except HandoffError as e:
        print(e)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
