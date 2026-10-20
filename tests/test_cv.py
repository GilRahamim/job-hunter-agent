import os

import pytest

from jobhunter import cv
from jobhunter.llm import get_llm
from jobhunter.skills import extract_skills

CV_TEXT = """Dana Levi
Junior Software Engineer
Skills: Python, C, SQL and Go. Built a RAG chatbot with LangChain and ChromaDB.
עבדתי בפייתון ובדוקר
"""


@pytest.fixture
def cfg(tmp_path):
    (tmp_path / "cv").mkdir()
    return {"profile": {"cv_dir": str(tmp_path / "cv"), "overrides": {}},
            "search": {"queries": ["AI Engineer", "Backend Developer"]},
            "llm": {"provider": "anthropic", "model": "x"}}


@pytest.fixture
def db_path(tmp_path):
    return tmp_path / "jobs.db"


class FakeLLM:
    def __init__(self, fail=False):
        self.calls = 0
        self.fail = fail

    def with_structured_output(self, model_cls):
        outer = self

        class S:
            def invoke(self, messages):
                outer.calls += 1
                if outer.fail:
                    raise RuntimeError("api down")
                return model_cls(name="Dana", skills=["Python", "RAG"], seniority="student",
                                 years_experience=0, target_roles=["AI Engineer"])
        return S()


def write_cv(cfg, name="cv.txt", text=CV_TEXT):
    path = cv.Path(cfg["profile"]["cv_dir"]) / name
    path.write_text(text, encoding="utf-8")
    return path


# ---------------------------------------------------------------- skills

@pytest.mark.parametrize("text,expected,absent", [
    ("Python, C, SQL and Go", {"Python", "C", "SQL", "Go"}, set()),
    ("I go to work and see a sea", set(), {"Go", "C"}),
    ("C++ and C# developer", {"C++", "C#"}, {"C"}),
    ("עבדתי בפייתון", {"Python"}, set()),
    ("R&D team", set(), {"R"}),
    ("Golang services", {"Go"}, set()),
])
def test_extract_skills(text, expected, absent):
    found = set(extract_skills(text))
    assert expected <= found
    assert not (absent & found)


# ---------------------------------------------------------------- reading

def test_find_cv_newest_wins_and_ignores_readme(cfg):
    folder = cv.Path(cfg["profile"]["cv_dir"])
    (folder / "README.md").write_text("ignore me")
    old = write_cv(cfg, "old.txt")
    os.utime(old, (1, 1))
    new = write_cv(cfg, "new.txt")
    assert cv.find_cv(folder) == new


def test_find_cv_missing_raises_clear_error(cfg):
    with pytest.raises(cv.CVError, match="No CV found"):
        cv.find_cv(cfg["profile"]["cv_dir"])


def test_read_docx_includes_tables(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Dana Levi")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python, Docker"
    path = tmp_path / "cv.docx"
    doc.save(path)
    text = cv.read_text(path)
    assert "Dana Levi" in text and "Python, Docker" in text


def test_scanned_pdf_gives_clear_error(tmp_path, monkeypatch):
    class Page:
        def extract_text(self):
            return ""

    class Reader:
        def __init__(self, path):
            self.pages = [Page()]

    monkeypatch.setattr("pypdf.PdfReader", Reader)
    pdf = tmp_path / "scan.pdf"
    pdf.write_bytes(b"%PDF")
    with pytest.raises(cv.CVError, match="text PDF or DOCX"):
        cv.read_text(pdf)


# ---------------------------------------------------------------- profile

def test_heuristic_profile_without_llm(cfg, db_path):
    write_cv(cfg)
    profile, _, hit = cv.load_profile(cfg, None, db_path)
    assert not hit
    assert profile.seniority == "junior" and profile.years_experience == 0
    assert {"Python", "SQL", "LangChain", "ChromaDB"} <= set(profile.skills)
    assert profile.target_roles == ["AI Engineer", "Backend Developer"]
    assert profile.name == "Dana Levi"


def test_llm_profile_is_cached_by_file_hash(cfg, db_path):
    write_cv(cfg)
    llm = FakeLLM()
    first, h1, hit1 = cv.load_profile(cfg, llm, db_path)
    second, h2, hit2 = cv.load_profile(cfg, llm, db_path)
    assert (hit1, hit2) == (False, True) and llm.calls == 1 and h1 == h2
    assert first == second and first.skills == ["Python", "RAG"]


def test_changed_cv_invalidates_cache(cfg, db_path):
    write_cv(cfg)
    llm = FakeLLM()
    _, h1, _ = cv.load_profile(cfg, llm, db_path)
    write_cv(cfg, text=CV_TEXT + "\nNew job at Acme")
    _, h2, hit = cv.load_profile(cfg, llm, db_path)
    assert h1 != h2 and not hit and llm.calls == 2


def test_llm_failure_falls_back_and_does_not_poison_cache(cfg, db_path):
    write_cv(cfg)
    profile, _, _ = cv.load_profile(cfg, FakeLLM(fail=True), db_path)
    assert "Python" in profile.skills and profile.seniority == "junior"
    good = FakeLLM()
    again, _, hit = cv.load_profile(cfg, good, db_path)     # later run with a working LLM
    assert not hit and good.calls == 1 and again.seniority == "student"


def test_overrides_applied_after_cache(cfg, db_path):
    write_cv(cfg)
    cv.load_profile(cfg, FakeLLM(), db_path)
    cfg["profile"]["overrides"] = {"add_skills": ["FastAPI"], "remove_skills": ["rag"], "seniority": "junior"}
    profile, _, hit = cv.load_profile(cfg, FakeLLM(), db_path)
    assert hit and profile.skills == ["Python", "FastAPI"] and profile.seniority == "junior"


# ---------------------------------------------------------------- llm.py

def test_get_llm_none_when_disabled_or_no_key(monkeypatch):
    monkeypatch.setattr("jobhunter.llm.load_env", lambda: None)   # ignore a real .env
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    assert get_llm({"llm": {"provider": "none", "model": "x"}}) is None
    assert get_llm({"llm": {"provider": "anthropic", "model": "x"}}) is None


def test_find_cv_skips_the_tailoring_template(cfg):
    folder = cv.Path(cfg["profile"]["cv_dir"])
    real = write_cv(cfg, "resume.txt")
    os.utime(real, (1, 1))
    (folder / "template.docx").write_bytes(b"")        # newer, but it is the tailoring template
    assert cv.find_cv(folder) == real
