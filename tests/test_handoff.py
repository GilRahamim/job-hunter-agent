import pytest

from jobhunter import db, handoff
from jobhunter.config import ROOT
from jobhunter.models import Job
from jobhunter.users import create_user

CV = "Dana Levi\ndana@example.com\nJunior Python Developer\nBuilt a RAG chatbot with LangChain.\n"
JOB = Job("greenhouse", "Junior Python Developer", "Acme Ltd", "https://acme/jobs/1", location="Tel Aviv, Israel",
          description="We build APIs in Python and FastAPI. Docker is an advantage.")


@pytest.fixture
def user():
    """A user with a CV and one scored job, all inside the temp folder."""
    user = create_user("Dana", "cv.txt", CV.encode())
    row = JOB.to_dict()
    row.update(score=81, prescore=55, verdict="strong", scored_by="llm:test", reasons="Python fits",
               matching_skills=["Python"], missing_skills=["FastAPI"], cv_tip="Lead with the RAG project")
    db.upsert_jobs([row], db.db_path(user=user))
    db.save_profile("h", "cv.txt", {"name": "Dana Levi"}, db.db_path(user=user))
    return user


def test_brief_holds_job_analysis_cv_and_rules(user, tmp_path):
    brief = handoff.write_brief(JOB.id, user)
    text = brief.read_text(encoding="utf-8")
    assert brief.parent == tmp_path / "users" / user / "data" / "tailored" / f"acme-ltd-junior-python-developer-{JOB.id[:6]}"
    assert "# Application kit: Junior Python Developer at Acme Ltd" in text
    assert "You are helping Dana Levi apply" in text
    assert "FastAPI. Docker is an advantage." in text                       # the full posting
    assert "Missing skills: FastAPI" in text and "CV tip: Lead with the RAG project" in text
    assert "Built a RAG chatbot with LangChain." in text                     # the CV itself
    assert "Never add a skill, tool, employer" in text                       # the no-fabrication rule
    assert str(brief.parent) in text and "--docx" in text and "cv_tailored.md" in text
    assert "looks English" in text and "written in Hebrew" in text           # posting vs. prep language


def test_brief_for_heuristic_job_and_unknown_id(user):
    row = Job("remotive", "דרוש/ה מפתח/ת פייתון", "", "https://x/2", description="ניסיון בפייתון").to_dict()
    row.update(score=40, prescore=40, scored_by="heuristic")
    db.upsert_jobs([row], db.db_path(user=user))
    text = handoff.write_brief(row["id"], user).read_text(encoding="utf-8")
    assert "keyword pre-score only" in text and "looks Hebrew" in text
    with pytest.raises(handoff.HandoffError):
        handoff.write_brief("nope", user)


def test_master_cv_is_included_when_present(user, tmp_path):
    (tmp_path / "users" / user / "cv" / "cv_master.yaml").write_text("projects: [rag]\n", encoding="utf-8")
    text = handoff.write_brief(JOB.id, user).read_text(encoding="utf-8")
    assert "## Master CV" in text and "projects: [rag]" in text


def test_prompt_is_short_and_relative_inside_the_project():
    brief = ROOT / "data" / "tailored" / "acme-x-123456" / "brief.md"
    assert handoff.claude_prompt(brief) == "Read data/tailored/acme-x-123456/brief.md and follow its instructions."
    assert handoff.shell_command(brief).startswith('claude "Read data/')


def test_subscription_env_drops_api_keys():
    env = handoff.subscription_env({"ANTHROPIC_API_KEY": "k", "ANTHROPIC_AUTH_TOKEN": "t", "PATH": "/bin"})
    assert env == {"PATH": "/bin"}


def test_launch_on_windows_opens_claude_without_the_api_key(user, monkeypatch):
    calls = []
    monkeypatch.setattr(handoff, "claude_path", lambda: r"C:\tools\claude.exe")
    monkeypatch.setattr(handoff.subprocess, "Popen", lambda cmd, **kw: calls.append((cmd, kw)))
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-secret")
    brief = handoff.write_brief(JOB.id, user)
    handoff.launch_claude(brief, system="windows")
    cmd, kw = calls[0]
    assert cmd == [r"C:\tools\claude.exe", handoff.claude_prompt(brief)]
    assert "ANTHROPIC_API_KEY" not in kw["env"] and kw["cwd"] == ROOT


def test_launch_on_macos_uses_terminal_and_unsets_keys(user, monkeypatch):
    calls = []
    monkeypatch.setattr(handoff, "claude_path", lambda: "/usr/local/bin/claude")
    monkeypatch.setattr(handoff.subprocess, "Popen", lambda cmd, **kw: calls.append(cmd))
    handoff.launch_claude(handoff.write_brief(JOB.id, user), system="darwin")
    script = calls[0][2]
    assert calls[0][0] == "osascript" and "unset ANTHROPIC_API_KEY ANTHROPIC_AUTH_TOKEN" in script


def test_launch_without_claude_or_elsewhere_explains(user, monkeypatch):
    brief = handoff.write_brief(JOB.id, user)
    monkeypatch.setattr(handoff, "claude_path", lambda: None)
    with pytest.raises(handoff.HandoffError, match="not installed"):
        handoff.launch_claude(brief, system="windows")
    assert not handoff.can_launch("windows")
    monkeypatch.setattr(handoff, "claude_path", lambda: "/usr/bin/claude")
    with pytest.raises(handoff.HandoffError, match="open a terminal"):
        handoff.launch_claude(brief, system="linux")
    assert handoff.can_launch("windows") and not handoff.can_launch("linux")


def test_markdown_to_docx(tmp_path):
    from docx import Document

    md = tmp_path / "cv_tailored.md"
    md.write_text("# Dana Levi\n\n## Skills\n- **Python**, SQL\n- LangChain\n\n---\nBuilt things.\n"
                  "ניסיון בפייתון\n", encoding="utf-8")
    out = handoff.markdown_to_docx(md)
    doc = Document(out)
    texts = [(p.style.name, p.text) for p in doc.paragraphs]
    assert texts[:4] == [("Heading 1", "Dana Levi"), ("Heading 2", "Skills"),
                         ("List Bullet", "Python, SQL"), ("List Bullet", "LangChain")]
    assert doc.paragraphs[2].runs[0].bold is True
    assert doc.paragraphs[-1].paragraph_format.alignment == 2           # Hebrew line is right-aligned


def test_kit_files_lists_what_claude_wrote(user):
    brief = handoff.write_brief(JOB.id, user)
    (brief.parent / "cover_letter.md").write_text("Dear Acme", encoding="utf-8")
    assert [p.name for p in handoff.kit_files(brief.parent)] == ["cover_letter.md"]


def test_cli_list_and_brief(user, capsys, monkeypatch):
    assert handoff.main(["--list", "--user", user]) == 0
    assert JOB.id in capsys.readouterr().out
    assert handoff.main([JOB.id, "--user", user]) == 0
    out = capsys.readouterr().out
    assert "brief:" in out and 'claude "Read' in out
    assert handoff.main(["missing", "--user", user]) == 1


# ------------------------------------------------------------ the owner's brief and Word edits

def _owner_cv(folder):
    """A small Word CV shaped like the owner's: name, title, contact, headings, entry, bullets, skills."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    for text, bold in [("Dana Levi", True), ("Software Engineer", None), ("GitHub | dana@example.com", None),
                       ("SUMMARY", True), ("Junior developer who builds Python services.", None),
                       ("EMPLOYMENT", True)]:
        doc.add_paragraph().add_run(text).bold = bold
    header = doc.add_paragraph()
    header.add_run("Developer, Acme.").bold = True
    header.add_run("\t2023 - Present")
    doc.add_paragraph("Built REST APIs in Python.", style="List Bullet")
    doc.add_paragraph()
    skills = doc.add_paragraph()
    skills.add_run("Deep Learning:").bold = True
    skills.add_run(" PyTorch, CNN")
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.size = Pt(10)
    folder.mkdir(parents=True, exist_ok=True)
    doc.save(folder / "Dana_Resume.docx")
    return folder / "Dana_Resume.docx"


def test_locked_paragraphs_follow_the_structure_lock(tmp_path):
    from docx import Document

    paragraphs = Document(_owner_cv(tmp_path / "cv")).paragraphs
    locked = [i for i, p in enumerate(paragraphs) if handoff._locked(i, p)]
    for i in (0, 2, 3, 5, 6, 8):          # name, contact, headings, entry header (tab), blank line
        assert i in locked, paragraphs[i].text
    for i in (1, 4, 7, 9):                # title line, summary, bullet, skills line
        assert i not in locked, paragraphs[i].text


def test_owner_brief_has_rules_notes_and_paragraph_map(user, tmp_path):
    cv_dir = tmp_path / "owner_cv"
    _owner_cv(cv_dir)
    (cv_dir / "notes").mkdir()
    (cv_dir / "notes" / "experience_bank.md").write_text("NOT: Kubernetes", encoding="utf-8")
    cfg = {"profile": {"cv_dir": str(cv_dir)}}
    folder = tmp_path / "kit"
    block = handoff.owner_block(cfg, folder)
    assert "owner-cv-rules.md" in block and "Structure lock" in block
    assert "NOT: Kubernetes" in block and "(not available)" in block       # master_cv.md is missing
    assert "  0 [locked] Dana Levi" in block and "  4 Junior developer" in block
    assert "<tab>" in block and "--apply-edits" in block and str(folder / "cv_edits.json") in block
    # other users keep the generic brief
    assert "owner's CV rules" not in handoff.write_brief(JOB.id, user).read_text(encoding="utf-8")
    assert handoff.owner_block({"profile": {"cv_dir": str(tmp_path / "empty")}}, folder) == ""


def test_apply_docx_edits_keeps_formatting_and_bolds(tmp_path):
    import json

    from docx import Document

    base = _owner_cv(tmp_path / "cv")
    edits = tmp_path / "kit" / "cv_edits.json"
    edits.parent.mkdir()
    edits.write_text(json.dumps({"out": "Dana_CV_Acme_Dev.docx", "edits": {
        "4": "Junior developer building **FastAPI** services in Python.",
        "9": "**Deep Learning:** CNN, PyTorch",
        "0": "Someone Else",
        "7": "Built REST APIs in Python and Docker for a team of eight engineers across two sites."}}),
        encoding="utf-8")
    out, warnings = handoff.apply_docx_edits(edits, base)
    doc = Document(out)
    assert out.name == "Dana_CV_Acme_Dev.docx" and len(doc.paragraphs) == len(Document(base).paragraphs)
    summary = doc.paragraphs[4]
    assert summary.text == "Junior developer building FastAPI services in Python."
    assert [r.bold for r in summary.runs] == [False, True, False]
    assert all(r.font.size.pt == 10 for r in summary.runs)                # the original run's formatting
    assert doc.paragraphs[9].runs[0].text == "Deep Learning:" and doc.paragraphs[9].runs[0].bold
    assert doc.paragraphs[0].text == "Dana Levi"                         # locked: untouched
    assert any("paragraph 0" in w for w in warnings)
    assert any("paragraph 7" in w and "wrap" in w for w in warnings)
    assert Document(base).paragraphs[4].text.startswith("Junior developer who")   # original unchanged


# ------------------------------------------------------------ headless

class FakeProc:
    def __init__(self, stdout="", returncode=0, stderr=""):
        self.stdout, self.returncode, self.stderr = stdout, returncode, stderr


def _fake_claude(monkeypatch, proc=None, exc=None, write=None):
    calls = []

    def run(cmd, **kw):
        calls.append((cmd, kw))
        if exc:
            raise exc
        if write:
            write()
        return proc

    monkeypatch.setattr(handoff, "claude_path", lambda: r"C:\tools\claude.exe")
    monkeypatch.setattr(handoff.subprocess, "run", run)
    return calls


def test_headless_command_stays_on_the_subscription():
    cmd = handoff.headless_command("claude", "abc", "dana")
    assert cmd[:3] == ["claude", "-p", "/apply abc --user dana --unattended"]
    assert "--bare" not in cmd
    assert cmd[cmd.index("--permission-mode") + 1] == "acceptEdits"
    assert cmd[cmd.index("--permission-prompts") + 1] == "none"
    assert cmd[cmd.index("--output-format") + 1] == "json"
    rules = [c for c in cmd if c.startswith("Bash(")]
    assert rules and all("-m jobhunter.handoff *)" in r for r in rules)
    assert handoff.headless_command("claude", "abc")[2] == "/apply abc --unattended"


def test_run_headless_drops_api_keys_and_logs(user, monkeypatch):
    import json

    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-secret")
    monkeypatch.setenv("ANTHROPIC_AUTH_TOKEN", "t")
    monkeypatch.setenv("CLAUDE_CODE_OAUTH_TOKEN", "oauth")
    folder = handoff.write_brief(JOB.id, user).parent
    out = json.dumps({"result": "Done: angle = APIs", "session_id": "s1", "total_cost_usd": 0.42})
    calls = _fake_claude(monkeypatch, FakeProc(out),
                         write=lambda: (folder / "cv_tailored.md").write_text("# Dana\n- Python", encoding="utf-8"))
    res = handoff.run_headless(JOB.id, user)
    cmd, kw = calls[0]
    assert res.ok and res.result == "Done: angle = APIs" and res.session_id == "s1" and res.folder == folder
    assert "ANTHROPIC_API_KEY" not in kw["env"] and "ANTHROPIC_AUTH_TOKEN" not in kw["env"]
    assert kw["env"]["CLAUDE_CODE_OAUTH_TOKEN"] == "oauth" and kw["cwd"] == ROOT and kw["timeout"] == 900
    assert "--user" in cmd[2] and "--bare" not in cmd
    log = json.loads((folder / "auto_tailor.json").read_text(encoding="utf-8"))
    assert log["exit_code"] == 0 and log["total_cost_usd"] == 0.42 and log["at"]
    names = [p.name for p in handoff.kit_files(folder)]
    assert "cv_tailored.docx" in names and "auto_tailor.json" not in names     # Word copy made in Python


def test_run_headless_usage_limit_stops_the_queue(user, monkeypatch):
    import json

    _fake_claude(monkeypatch, FakeProc(json.dumps({"is_error": True, "result": "Claude usage limit reached"}), 1))
    res = handoff.run_headless(JOB.id, user)
    assert not res.ok and res.limit_hit and res.exit_code == 1


def test_run_headless_failures_never_raise(user, monkeypatch):
    import subprocess

    _fake_claude(monkeypatch, exc=subprocess.TimeoutExpired("claude", 5))
    res = handoff.run_headless(JOB.id, user, timeout=5)
    assert not res.ok and not res.limit_hit and "timed out" in res.result
    assert (res.folder / "auto_tailor.json").exists()

    _fake_claude(monkeypatch, FakeProc("not json", 2, "boom"))
    res = handoff.run_headless(JOB.id, user)
    assert not res.ok and not res.limit_hit and res.exit_code == 2

    _fake_claude(monkeypatch, exc=OSError("cannot start"))
    assert "cannot start" in handoff.run_headless(JOB.id, user).result

    assert "no job" in handoff.run_headless("missing", user).result

    monkeypatch.setattr(handoff, "claude_path", lambda: None)
    res = handoff.run_headless(JOB.id, user)
    assert not res.ok and "not installed" in res.result


def test_cli_headless(user, capsys, monkeypatch):
    import json

    _fake_claude(monkeypatch, FakeProc(json.dumps({"result": "all written"})))
    assert handoff.main([JOB.id, "--user", user, "--headless"]) == 0
    assert "ok in" in capsys.readouterr().out
    _fake_claude(monkeypatch, FakeProc("", 1, "Login expired"))
    assert handoff.main([JOB.id, "--user", user, "--headless"]) == 1
